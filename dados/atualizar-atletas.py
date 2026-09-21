#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Atualiza a lista de atletas de campeonato-inicial.json a partir da planilha.

    python3 dados/atualizar-atletas.py ATLETAS_TOTAIS.xlsx            # so mostra
    python3 dados/atualizar-atletas.py ATLETAS_TOTAIS.xlsx --aplicar  # grava

A planilha tem uma categoria por coluna e, a partir da linha marcada
ATLETAS DESISTENTES, os nomes que sairam do campeonato — esses nao entram.

A planilha nao traz a unidade de cada atleta, entao quem ja estava cadastrado
precisa ser reconhecido para manter a unidade e o id (o id e o que amarra os
resultados ja lancados). Reconhecer nao pode ser por igualdade: os nomes vem
encurtados de uma planilha para a outra ("Heloisa Machado" onde antes havia
"Heloisa Machado Agostini") e a grafia varia (Thiago/Tiago, Victor/Vitor,
Cesar/Cezar). Dai as regras de casamento em pontuar(), da mais forte para a
mais fraca, e a atribuicao gulosa: todos os pares possiveis sao ordenados por
pontuacao e o melhor leva, para um nome curto nao roubar o atleta de outro so
por vir antes na planilha.

Quem nao for reconhecido entra como atleta novo, sem unidade — a unidade se
preenche na tela do organizador. Conferir sempre a secao CASARAM POR
APROXIMACAO antes de aplicar.

Depois de aplicar: commitar o JSON, publicar, e clicar em CARGA INICIAL no
organizador — e ela que leva a lista para o Firestore.
"""
import json, io, os, re, sys, unicodedata
from collections import Counter
from difflib import SequenceMatcher

try:
    import openpyxl
except ImportError:
    sys.exit('falta o openpyxl: pip install openpyxl')

AQUI = os.path.dirname(os.path.abspath(__file__))
JSON = os.path.join(AQUI, 'campeonato-inicial.json')
ABA = 'ATLETAS 2026'
MARCA_DESISTENTES = 'ATLETAS DESISTENTES'

CATEGORIA = {
    'ELITE MASCULINO': 'Elite Masculino',
    'ELITE FEMININO': 'Elite Feminino',
    'RX MASCULINO': 'RX Masculino',
    'RX FEMININO': 'RX Feminino',
    'INTERMEDIARIO MASCULINO': 'Intermediário Masculino',
    'INTERMEDIARIO FEMININO': 'Intermediário Feminino',
    'SCALED MASCULINO': 'Scaled Masculino',
    'SCALED FEMININO': 'Scaled Feminino',
    'MASTER MASCULINO': 'Master 45+ Masculino',
    'MASTER FEMININO': 'Master 45+ Feminino',
}

# sobrenomes comuns demais para identificar alguem sozinhos
COMUNS = {'silva', 'santos', 'souza', 'sousa', 'oliveira', 'lima', 'costa', 'pereira',
          'ferreira', 'alves', 'rodrigues', 'martins', 'carvalho', 'gomes', 'ribeiro',
          'almeida', 'araujo', 'barbosa', 'rocha', 'dias', 'nunes', 'moraes', 'morais'}


def sem_acento(s):
    return ''.join(c for c in unicodedata.normalize('NFKD', s) if not unicodedata.combining(c))


def normal(nome):
    s = re.sub(r'\(.*?\)', ' ', sem_acento(nome))
    return ' '.join(re.sub(r'[^a-zA-Z ]', ' ', s).lower().split())


def fonetico(nome):
    """normal(), mais as trocas de grafia que aparecem na pratica."""
    s = normal(nome)
    s = s.replace('ph', 'f').replace('th', 't').replace('ct', 't')
    s = s.replace('y', 'i').replace('z', 's').replace('h', '')
    s = re.sub(r'(.)\1+', r'\1', s)   # letras dobradas: Cappellano -> Capelano
    s = re.sub(r'm\b', 'n', s)        # William -> Willian
    return ' '.join(s.split())


def apelido(nome):
    """O que estiver entre parenteses, normalizado; '' se nao houver."""
    m = re.search(r'\((.*?)\)', nome)
    return normal(m.group(1)) if m else ''


def feminina(categoria):
    return categoria.endswith('Feminino')


def limpar(v):
    return ' '.join(str(v).split()) if v and str(v).strip() else ''


def ler_planilha(caminho):
    """Devolve (ativos, desistentes), cada um como lista de (nome, categoria)."""
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = wb[ABA] if ABA in wb.sheetnames else wb[wb.sheetnames[0]]
    linhas = list(ws.iter_rows(values_only=True))
    if not linhas:
        sys.exit('planilha vazia')
    cabec = [limpar(h) for h in linhas[0]]
    desconhecidas = [c for c in cabec if c and c not in CATEGORIA]
    if desconhecidas:
        sys.exit('coluna sem categoria conhecida: %s' % desconhecidas)

    corte = None
    for i, linha in enumerate(linhas[1:], start=1):
        if any(limpar(v).upper() == MARCA_DESISTENTES for v in linha):
            corte = i
            break
    if corte is None:
        print('aviso: nao achei a linha "%s" — tratando a planilha inteira como ativos'
              % MARCA_DESISTENTES)
        corte = len(linhas)

    def colher(bloco):
        fora = []
        for linha in bloco:
            for col, valor in zip(cabec, linha):
                nome = limpar(valor)
                if nome and nome.upper() != MARCA_DESISTENTES:
                    fora.append((nome, CATEGORIA[col]))
        return fora

    return colher(linhas[1:corte]), colher(linhas[corte + 1:])


def pontuar(nome, categoria, antigo):
    """0 a 100: o quanto este nome da planilha e este atleta ja cadastrado sao a
    mesma pessoa. 0 = nao sao."""
    n, fn = normal(nome), fonetico(nome)
    na, fna = normal(antigo['nome']), fonetico(antigo['nome'])
    if na == n:
        return 100, 'nome igual'

    # generos diferentes so casam por nome igual (que e como uma correcao de
    # categoria aparece), e esse caso ja passou acima
    if feminina(categoria) != feminina(antigo['categoria']):
        return 0, ''

    ap = apelido(nome)
    if ap and (ap == na or ap in set(na.split())):
        return 97, 'apelido "%s"' % ap

    tfn, tfa = set(fn.split()), set(fna.split())
    if tfn and tfa and (tfn <= tfa or tfa <= tfn) and fn.split()[0] == fna.split()[0]:
        return 95 - abs(len(tfn) - len(tfa)), 'um nome e a versao curta do outro'

    # parecidos no todo, desde que primeiro nome e ultimo sobrenome batam: sem a
    # exigencia do sobrenome, "Stephanie Martins" casaria com "Stefhani Mattos"
    u, ua = fn.split()[-1], fna.split()[-1]
    mesmo_sobrenome = (SequenceMatcher(None, u, ua).ratio() >= 0.8
                       or (min(len(u), len(ua)) >= 3 and (u.startswith(ua) or ua.startswith(u))))
    r = SequenceMatcher(None, fn, fna).ratio()
    if r >= 0.82 and fn.split()[0][:3] == fna.split()[0][:3] and mesmo_sobrenome:
        return r * 100, 'parecidos (%.0f%%)' % (r * 100)

    # sobrenome incomum igual, mesma categoria: pega quem passou a usar o apelido
    # no lugar do nome ("Tato Outor" / "Luiz Felipe Outor"). Nunca o primeiro
    # nome, senao todo "Joao" casaria com todo "Joao".
    if antigo['categoria'] == categoria:
        def sobrenomes(toks):
            return {t for t in toks[1:] if len(t) >= 4 and t not in COMUNS}
        if sobrenomes(fn.split()) & sobrenomes(fna.split()):
            return 70, 'mesmo sobrenome incomum, mesma categoria'
    return 0, ''


def casar(ativos, antigos):
    """Liga cada nome da planilha a um atleta ja cadastrado, do par mais parecido
    para o menos. Devolve {indice na planilha: (antigo, pontuacao, motivo)}."""
    pares = []
    for i, (nome, cat) in enumerate(ativos):
        for a in antigos:
            p, motivo = pontuar(nome, cat, a)
            if p:
                pares.append((p, i, a['id'], motivo))
    pares.sort(key=lambda x: (-x[0], x[1], x[2]))

    por_id = {a['id']: a for a in antigos}
    escolha, tomados = {}, set()
    for p, i, aid, motivo in pares:
        if i not in escolha and aid not in tomados:
            escolha[i] = (por_id[aid], p, motivo)
            tomados.add(aid)
    return escolha


def slug(nome):
    s = re.sub(r'\(.*?\)', ' ', sem_acento(nome))
    return re.sub(r'[^a-zA-Z0-9]+', '-', s).strip('-').lower() or 'atleta'


def montar(ativos, escolha, categorias):
    saida, ids = [], set()
    for i, (nome, cat) in enumerate(ativos):
        antigo = escolha.get(i, (None,))[0]
        if antigo:
            aid, unidade = antigo['id'], antigo.get('unidade', '')
        else:
            base, n = 'at_' + slug(nome), 2
            aid = base
            while aid in ids:
                aid, n = '%s-%d' % (base, n), n + 1
            unidade = ''
        ids.add(aid)
        saida.append({'id': aid, 'nome': nome, 'categoria': cat, 'unidade': unidade})

    ordem = {c: i for i, c in enumerate(categorias)}
    saida.sort(key=lambda a: (ordem.get(a['categoria'], 99), a['nome'].lower()))
    return saida


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    caminho = sys.argv[1]
    aplicar = '--aplicar' in sys.argv[2:]

    ativos, desistentes = ler_planilha(caminho)
    dados = json.load(io.open(JSON, encoding='utf-8'))
    antigos = dados['atletas']
    escolha = casar(ativos, antigos)
    saida = montar(ativos, escolha, dados['categorias'])

    print('planilha: %d ativos, %d desistentes | cadastro atual: %d atletas\n'
          % (len(ativos), len(desistentes), len(antigos)))

    aprox = [(p, ativos[i][0], a, motivo)
             for i, (a, p, motivo) in escolha.items() if p < 97]
    print('CASARAM POR APROXIMACAO (%d) — conferir um a um:' % len(aprox))
    for p, nome, a, motivo in sorted(aprox):
        print('  %3.0f  %-28s ~ %-32s [%s]  %s'
              % (p, nome, a['nome'], a.get('unidade', ''), motivo))

    muda = [(nome, escolha[i][0], cat) for i, (nome, cat) in enumerate(ativos)
            if i in escolha and escolha[i][0]['categoria'] != cat]
    print('\nMUDAM DE CATEGORIA (%d):' % len(muda))
    for nome, a, cat in sorted(muda):
        print('  %-28s %-24s -> %s' % (nome, a['categoria'], cat))

    novos = [(nome, cat) for i, (nome, cat) in enumerate(ativos) if i not in escolha]
    print('\nENTRAM, sem unidade (%d):' % len(novos))
    for nome, cat in sorted(novos, key=lambda x: (x[1], x[0])):
        print('  %-28s %s' % (nome, cat))

    ficam = {a['id'] for a, _, _ in escolha.values()}
    saem = [a for a in antigos if a['id'] not in ficam]
    desist = {normal(n) for n, _ in desistentes}
    print('\nSAEM (%d):' % len(saem))
    for a in sorted(saem, key=lambda x: (x['categoria'], x['nome'])):
        na = normal(a['nome'])
        marcado = any(d == na or d in na or na in d for d in desist)
        print('  %-32s %-24s%s' % (a['nome'], a['categoria'],
                                   '  (na lista de desistentes)' if marcado else ''))

    repetidos = [k for k, v in Counter((a['nome'].lower(), a['categoria'])
                                       for a in saida).items() if v > 1]
    if repetidos:
        print('\nATENCAO — mesmo nome na mesma categoria: %s' % repetidos)

    print('\nRESULTADO: %d atletas' % len(saida))
    c = Counter(a['categoria'] for a in saida)
    for k in dados['categorias']:
        print('  %-26s %d' % (k, c[k]))

    if not aplicar:
        print('\n(nada foi gravado — rode de novo com --aplicar)')
        return

    assert len(saida) == len(ativos), 'perdi atletas no caminho'
    assert len({a['id'] for a in saida}) == len(saida), 'id repetido'
    dados['atletas'] = saida
    # anexos chegam com um prefixo de hash no nome; guardar so o nome da planilha
    planilha = re.sub(r'^[0-9a-f]{6,}-', '', os.path.basename(caminho))
    dados['_comentario'] = ('Lista oficial do Blacksheep Invitational, gerada de %s '
                            'por dados/atualizar-atletas.py.' % planilha)
    io.open(JSON, 'w', encoding='utf-8').write(
        json.dumps(dados, ensure_ascii=False, indent=2) + '\n')
    print('\ngravado em %s' % JSON)
    print('agora: commitar, publicar, e clicar em CARGA INICIAL no organizador.')


if __name__ == '__main__':
    main()
