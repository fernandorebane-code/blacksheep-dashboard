#!/usr/bin/env python3
"""Monta a planilha de baterias a partir da classificacao exportada do leaderboard.

Entrada: o CSV do botao "CSV DE TODAS" do leaderboard
         (Categoria; Pos; Atleta; Unidade; Total; Provas valendo; Com resultado;
          Completo; Faltando).

Regras:
  - So entra quem esta COMPLETO, isto e, quem tem resultado em todas as provas
    que ja estao valendo. Quem ainda deve prova nao vai para a bateria.
  - Dentro da categoria a ordem e da PIOR colocacao para a MELHOR: a ultima
    bateria e a dos lideres.
  - Baterias de RAIAS atletas. Quando a categoria nao fecha um multiplo exato,
    a bateria menor e a PRIMEIRA (a dos piores), para a bateria final sair cheia.

Uso:
    python3 dados/montar-baterias.py classificacao_todas_categorias.csv
    python3 dados/montar-baterias.py entrada.csv -o baterias-domingo.xlsx
"""
import argparse, csv, os, re, sys, unicodedata
from collections import OrderedDict
from datetime import datetime, timedelta

# ------------------------------------------------------------------ config
TITULO   = 'PROVA 4 E 5'      # vai no topo de cada bloco, junto do nome da trilha
RAIAS    = 8                  # atletas por bateria
INICIO   = '08:00'
INTERVALO = 18                # minutos entre uma bateria e a seguinte, na mesma trilha
ABA      = 'BATERIAS DOMINGO'

# Cada trilha vira uma coluna de blocos na planilha, com o proprio relogio.
# A ordem das categorias aqui e a ordem em que elas entram na trilha.
TRILHAS = [
    ('SCALED',                  ['Scaled Feminino', 'Scaled Masculino']),
    ('INTERMEDIARIO E MASTER',  ['Intermediário Feminino', 'Master 45+ Feminino',
                                 'Intermediário Masculino', 'Master 45+ Masculino']),
    ('RX',                      ['RX Feminino', 'RX Masculino']),
    ('ELITE',                   ['Elite Feminino', 'Elite Masculino']),
]

# ------------------------------------------------------------------ leitura
def chave(txt):
    """Normaliza para comparar nome de categoria com nome de arquivo.

    O CSV por categoria nasce com o nome da categoria no arquivo, e no caminho
    ele chega sem acento e com a pontuacao trocada: "Intermediário Masculino"
    vira "intermedia_rio_masculino". Tirando acento e tudo que nao e letra ou
    numero, os dois viram "intermediariomasculino".
    """
    sem = unicodedata.normalize('NFKD', txt)
    sem = ''.join(c for c in sem if not unicodedata.combining(c))
    return re.sub(r'[^a-z0-9]', '', sem.lower())


def categoria_do_arquivo(caminho, conhecidas):
    k = chave(os.path.basename(caminho))
    achadas = [c for c in conhecidas if chave(c) and chave(c) in k]
    if not achadas:
        return None
    return max(achadas, key=lambda c: len(chave(c)))   # a mais especifica


def ler_por_categoria(caminhos, conhecidas):
    """CSV do botao CSV (uma categoria por arquivo, com as provas em colunas).

    Aqui nao existe coluna "Completo": a prova conta como valendo quando ao menos
    um atleta do arquivo tem resultado nela, e o atleta esta completo quando tem
    resultado em todas essas.
    """
    saida = []
    for caminho in caminhos:
        cat = categoria_do_arquivo(caminho, conhecidas)
        if not cat:
            sys.exit('Nao consegui dizer de que categoria e o arquivo ' +
                     os.path.basename(caminho) + '. Renomeie com o nome da categoria ' +
                     'ou use o CSV DE TODAS.')
        linhas = ler(caminho)
        provas = [k for k in (linhas[0] if linhas else {}) if k.strip().endswith('(result)')]
        valendo = [p for p in provas if any((l.get(p) or '').strip() for l in linhas)]
        for l in linhas:
            faltando = [p[:-len(' (result)')].strip() for p in valendo
                        if not (l.get(p) or '').strip()]
            saida.append({
                'Categoria': cat, 'Pos': coluna(l, 'Pos'),
                'Atleta': coluna(l, 'Atleta', 'Nome'), 'Unidade': coluna(l, 'Unidade'),
                'Total': coluna(l, 'Total'),
                'Completo': 'nao' if faltando else 'sim',
                'Faltando': ' | '.join(faltando),
            })
    return saida


def ler(caminho):
    with open(caminho, encoding='utf-8-sig', newline='') as f:
        amostra = f.read(4096); f.seek(0)
        sep = ';' if amostra.count(';') >= amostra.count(',') else ','
        return list(csv.DictReader(f, delimiter=sep))


def coluna(linha, *nomes):
    for n in nomes:
        for k in linha:
            if k.strip().lower() == n.lower():
                return (linha[k] or '').strip()
    return ''


def baterias_da_categoria(atletas, raias):
    """Da pior para a melhor colocacao, com a bateria menor na frente."""
    ordenado = sorted(atletas, key=lambda a: -a['pos'])       # pior primeiro
    n = len(ordenado)
    resto = n % raias
    blocos, i = [], 0
    if resto:                                                 # a menor vem primeiro
        blocos.append(ordenado[:resto]); i = resto
    while i < n:
        blocos.append(ordenado[i:i + raias]); i += raias
    return blocos


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('csv', nargs='+',
                    help='o CSV DE TODAS, ou varios CSV por categoria')
    ap.add_argument('-o', '--saida', default='baterias.xlsx')
    ap.add_argument('--raias', type=int, default=RAIAS)
    ap.add_argument('--inicio', default=INICIO)
    ap.add_argument('--intervalo', type=int, default=INTERVALO)
    ap.add_argument('--titulo', default=TITULO)
    ap.add_argument('--aba', default=ABA)
    args = ap.parse_args()

    conhecidas_cfg = [c for _, cats in TRILHAS for c in cats]
    primeiro = ler(args.csv[0])
    if not primeiro:
        sys.exit('CSV vazio.')
    if any(k.strip().lower() == 'categoria' for k in primeiro[0]):
        if len(args.csv) > 1:
            sys.exit('O CSV DE TODAS ja traz tudo — passe um arquivo so.')
        linhas = primeiro
    else:
        linhas = ler_por_categoria(args.csv, conhecidas_cfg)

    porCat, fora, semPos = OrderedDict(), [], []
    for l in linhas:
        cat = coluna(l, 'Categoria')
        nome = coluna(l, 'Atleta', 'Nome')
        if not cat or not nome:
            continue
        if coluna(l, 'Completo').lower() not in ('sim', 's', 'yes', '1'):
            fora.append((cat, nome, coluna(l, 'Faltando'))); continue
        pos = coluna(l, 'Pos')
        if not pos.isdigit():
            semPos.append((cat, nome)); continue
        porCat.setdefault(cat, []).append(
            {'nome': nome, 'unidade': coluna(l, 'Unidade'), 'pos': int(pos)})

    conhecidas = set(conhecidas_cfg)
    orfas = [c for c in porCat if c not in conhecidas]
    ausentes = [c for c in conhecidas_cfg if c not in porCat]

    # ---------------------------------------------------------- planilha
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook(); ws = wb.active; ws.title = args.aba
    fina = Side(style='thin', color='000000')
    borda = Border(left=fina, right=fina, top=fina, bottom=fina)
    tit = Font(bold=True, italic=True, underline='single', size=12)
    cab = Font(bold=True, italic=True, underline='single', size=11)
    meio = Alignment(horizontal='center', vertical='center', wrap_text=True)

    resumo = []
    for t, (trilha, cats) in enumerate(TRILHAS):
        c0 = 1 + t * 5                                   # A, F, K, P...
        relogio = datetime.strptime(args.inicio, '%H:%M')
        linha = 1
        for cat in cats:
            for bloco in baterias_da_categoria(porCat.get(cat, []), args.raias):
                hora = relogio.strftime('%H:%M')
                ws.merge_cells(start_row=linha, start_column=c0,
                               end_row=linha, end_column=c0 + 3)
                cel = ws.cell(linha, c0, f'{args.titulo} {trilha}')
                cel.font = tit; cel.alignment = meio
                for j, txt in enumerate(['CATEGORIA', 'RAIA:', 'NOME:', 'HORÁRIO:']):
                    c = ws.cell(linha + 1, c0 + j, txt)
                    c.font = cab; c.alignment = meio; c.border = borda
                p, u = linha + 2, linha + 1 + args.raias
                ws.merge_cells(start_row=p, start_column=c0, end_row=u, end_column=c0)
                cc = ws.cell(p, c0, cat.upper()); cc.alignment = meio; cc.border = borda
                ws.merge_cells(start_row=p, start_column=c0 + 3, end_row=u, end_column=c0 + 3)
                ch = ws.cell(p, c0 + 3, hora); ch.alignment = meio; ch.border = borda
                for r in range(args.raias):
                    atleta = bloco[r] if r < len(bloco) else None
                    cr = ws.cell(p + r, c0 + 1, r + 1); cr.alignment = meio; cr.border = borda
                    cn = ws.cell(p + r, c0 + 2, atleta['nome'] if atleta else '')
                    cn.alignment = Alignment(vertical='center'); cn.border = borda
                    for cx in (c0, c0 + 3):
                        ws.cell(p + r, cx).border = borda
                resumo.append((trilha, cat, hora, [a['nome'] for a in bloco]))
                linha = u + 1
                relogio += timedelta(minutes=args.intervalo)
        # get_column_letter, e nao ws.cell(...): a linha 1 esta mesclada e a
        # celula mesclada nao carrega column_letter.
        for desloc, larg in ((0, 26), (1, 7), (2, 28), (3, 11), (4, 3)):
            if desloc == 4 and t == len(TRILHAS) - 1:
                continue
            ws.column_dimensions[get_column_letter(c0 + desloc)].width = larg

    wb.save(args.saida)

    # ---------------------------------------------------------- relatorio
    print(f'{args.saida}  ({args.aba})\n')
    for trilha, cat, hora, nomes in resumo:
        print(f'  {hora}  {trilha:24s} {cat:26s} {len(nomes)} atleta(s)')
        for i, n in enumerate(nomes, 1):
            print(f'           raia {i}: {n}')
    print()
    total = sum(len(n) for _, _, _, n in resumo)
    print(f'{total} atleta(s) escalado(s) em {len(resumo)} bateria(s).')
    if fora:
        print(f'\nFORA — ainda devem prova ({len(fora)}):')
        for cat, nome, falta in fora:
            print(f'  {cat:26s} {nome:28s} falta: {falta or "(sem resultado nenhum)"}')
    if semPos:
        print(f'\nFORA — sem colocação ({len(semPos)}):')
        for cat, nome in semPos:
            print(f'  {cat:26s} {nome}')
    if orfas:
        print('\nATENÇÃO — categorias que não estão em nenhuma trilha do script:')
        for c in orfas:
            print(f'  {c}  ({len(porCat[c])} atleta(s) COMPLETOS ficaram de fora)')
    if ausentes:
        print('\nSEM DADOS — categorias da configuração que não vieram em nenhum CSV:')
        for c in ausentes:
            print(f'  {c}')


if __name__ == '__main__':
    main()
